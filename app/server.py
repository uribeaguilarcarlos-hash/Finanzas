import os
import re
import glob
import json
import shutil
import sqlite3
import tempfile
import collections
from datetime import datetime, date
from pathlib import Path
from typing import Optional, List, Dict, Any

import openpyxl
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel

BASE_DIR = Path(__file__).resolve().parent
WORKSPACE_DIR = BASE_DIR.parent
DB_PATH = BASE_DIR / "ml_statements.db"
STATIC_DIR = BASE_DIR / "static"
CONCENTRADO_PATH = WORKSPACE_DIR / "CONCENTRADO ML LOGISTICS ELIEZER 2026.xlsx"

app = FastAPI(title="ML Logistics - Estados de Cuenta & Statements Vertical")


def get_db():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    cur = conn.cursor()
    cur.executescript("""
        CREATE TABLE IF NOT EXISTS statements (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_date TEXT UNIQUE NOT NULL,
            stmt_date TEXT NOT NULL,
            filename TEXT NOT NULL,
            file_size INTEGER DEFAULT 0,
            file_mtime REAL DEFAULT 0,
            imported_at TEXT NOT NULL,
            total_containers INTEGER DEFAULT 0,
            total_cost_master REAL DEFAULT 0,
            total_previo REAL DEFAULT 0,
            total_insurance REAL DEFAULT 0,
            acc_balance REAL DEFAULT 0,
            acc_pending_receive REAL DEFAULT 0,
            acc_pending_note TEXT DEFAULT '',
            acc_total_rcvd REAL DEFAULT 0,
            acc_total_cost REAL DEFAULT 0,
            acc_paid_mexico REAL DEFAULT 0,
            acc_paid_china REAL DEFAULT 0,
            sheets_json TEXT DEFAULT '[]'
        );

        CREATE TABLE IF NOT EXISTS master_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_date TEXT NOT NULL,
            row_num INTEGER NOT NULL,
            ref_key TEXT NOT NULL,
            ref_no TEXT NOT NULL,
            customer TEXT DEFAULT '',
            customer_for TEXT DEFAULT '',
            container TEXT DEFAULT '',
            etd TEXT DEFAULT '',
            eta TEXT DEFAULT '',
            update_eta TEXT DEFAULT '',
            shipping_line TEXT DEFAULT '',
            pol TEXT DEFAULT '',
            pod TEXT DEFAULT '',
            bill_amt REAL,
            bill_amt_text TEXT DEFAULT '',
            sea_freight REAL DEFAULT 0,
            fob_cost REAL DEFAULT 0,
            amend_fee REAL DEFAULT 0,
            exw_fee REAL DEFAULT 0,
            inland_fee REAL DEFAULT 0,
            total_cost REAL DEFAULT 0,
            ttl_cost_profit REAL,
            ttl_cost_profit_text TEXT DEFAULT '',
            previo REAL DEFAULT 0,
            insurance REAL DEFAULT 0,
            remark TEXT DEFAULT '',
            UNIQUE(raw_date, ref_key)
        );

        CREATE INDEX IF NOT EXISTS idx_master_raw_date ON master_items(raw_date);
        CREATE INDEX IF NOT EXISTS idx_master_ref_no ON master_items(ref_no);

        CREATE TABLE IF NOT EXISTS accounting_payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_date TEXT NOT NULL,
            section TEXT NOT NULL,
            row_num INTEGER NOT NULL,
            rcvd_date TEXT DEFAULT '',
            rcvd_amt REAL DEFAULT 0,
            rcvd_amt_text TEXT DEFAULT '',
            paid_from TEXT DEFAULT '',
            paid_to TEXT DEFAULT '',
            remark TEXT DEFAULT '',
            payment_sig TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_acc_pay_date ON accounting_payments(raw_date);

        CREATE TABLE IF NOT EXISTS accounting_customer_costs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            raw_date TEXT NOT NULL,
            row_num INTEGER NOT NULL,
            customer_name TEXT NOT NULL,
            cost REAL DEFAULT 0,
            remark TEXT DEFAULT ''
        );

        CREATE INDEX IF NOT EXISTS idx_acc_cust_date ON accounting_customer_costs(raw_date);

        CREATE TABLE IF NOT EXISTS review_status (
            change_key TEXT PRIMARY KEY,
            raw_date TEXT NOT NULL,
            change_type TEXT NOT NULL,
            ref_no TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            note TEXT DEFAULT '',
            updated_at TEXT NOT NULL
        );

        CREATE INDEX IF NOT EXISTS idx_review_date ON review_status(raw_date);
    """)
    conn.commit()
    conn.close()


def parse_float(val: Any) -> Optional[float]:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return round(float(val), 2)
    s = str(val).strip()
    if not s or s == "-":
        return None
    cleaned = re.sub(r"[,$\s]", "", s)
    try:
        return round(float(cleaned), 2)
    except ValueError:
        return None


def parse_date_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return ""
    if " 00:00:00" in s:
        s = s.replace(" 00:00:00", "")
    return s[:25]


def safe_load_workbook(filepath: Path, data_only: bool = True, read_only: bool = True):
    try:
        return openpyxl.load_workbook(str(filepath), data_only=data_only, read_only=read_only), None
    except PermissionError:
        tmp_path = Path(tempfile.gettempdir()) / f"ml_tmp_{os.getpid()}_{filepath.name}"
        shutil.copy2(str(filepath), str(tmp_path))
        return openpyxl.load_workbook(str(tmp_path), data_only=data_only, read_only=read_only), tmp_path


def extract_raw_date_from_filename(filename: str) -> Optional[str]:
    m = re.search(r"(202\d{5})", filename)
    if m:
        return m.group(1)
    return None


def format_stmt_date(raw_date: str) -> str:
    if len(raw_date) == 8 and raw_date.isdigit():
        return f"{raw_date[:4]}-{raw_date[4:6]}-{raw_date[6:8]}"
    return raw_date


def parse_and_store_statement(filepath: Path, force: bool = False) -> Dict[str, Any]:
    raw_date = extract_raw_date_from_filename(filepath.name)
    if not raw_date:
        raise ValueError(f"No se encontró fecha AAAAMMDD en el archivo: {filepath.name}")

    stat = filepath.stat()
    conn = get_db()
    cur = conn.cursor()

    existing = cur.execute(
        "SELECT id, file_size, file_mtime FROM statements WHERE raw_date = ?",
        (raw_date,)
    ).fetchone()

    if existing and not force:
        if existing["file_size"] == stat.st_size and abs(existing["file_mtime"] - stat.st_mtime) < 1.0:
            conn.close()
            return {"raw_date": raw_date, "status": "cached"}

    wb, tmp_path = safe_load_workbook(filepath)
    try:
        sheetnames = wb.sheetnames
        ml_sheet_name = next((s for s in sheetnames if "master" in s.lower()), None)
        acc_sheet_name = next((s for s in sheetnames if "accounting" in s.lower()), None)

        if not ml_sheet_name:
            raise ValueError(f"El archivo {filepath.name} no contiene la hoja 'Master list'")

        ws_ml = wb[ml_sheet_name]
        ml_rows = list(ws_ml.iter_rows(values_only=True))

        header_idx = None
        for idx, r in enumerate(ml_rows[:15]):
            if r and any(str(c).strip().upper().startswith("REF") for c in r if c is not None):
                header_idx = idx
                break

        if header_idx is None:
            header_idx = 4

        master_records = []
        seen_keys = set()

        for idx in range(header_idx + 1, len(ml_rows)):
            r = ml_rows[idx]
            if not r or r[0] is None or str(r[0]).strip() == "":
                continue
            ref_no = str(r[0]).strip()
            if ref_no.upper().startswith("TOTAL") or ref_no.upper().startswith("REF/INV"):
                continue

            def cell(col_i):
                return r[col_i] if col_i < len(r) else None

            container = str(cell(3) or "").strip()
            ref_key = ref_no
            if ref_key in seen_keys:
                ref_key = f"{ref_no}__ROW{idx+1}"
            seen_keys.add(ref_key)

            bill_raw = cell(10)
            bill_num = parse_float(bill_raw)
            bill_txt = str(bill_raw).strip() if bill_raw is not None and bill_num is None else ""

            ttl_p_raw = cell(17)
            ttl_p_num = parse_float(ttl_p_raw)
            ttl_p_txt = str(ttl_p_raw).strip() if ttl_p_raw is not None and ttl_p_num is None else ""

            sea_freight = parse_float(cell(11)) or 0.0
            fob_cost = parse_float(cell(12)) or 0.0
            amend_fee = parse_float(cell(13)) or 0.0
            exw_fee = parse_float(cell(14)) or 0.0
            inland_fee = parse_float(cell(15)) or 0.0
            total_cost = parse_float(cell(16))
            if total_cost is None:
                total_cost = round(sea_freight + fob_cost + amend_fee + exw_fee + inland_fee, 2)

            previo = parse_float(cell(18)) or 0.0
            insurance = parse_float(cell(19)) or 0.0

            master_records.append((
                raw_date,
                idx + 1,
                ref_key,
                ref_no,
                str(cell(1) or "").strip(),
                str(cell(2) or "").strip(),
                container,
                parse_date_str(cell(4)),
                parse_date_str(cell(5)),
                parse_date_str(cell(6)),
                str(cell(7) or "").strip(),
                str(cell(8) or "").strip(),
                str(cell(9) or "").strip(),
                bill_num,
                bill_txt,
                sea_freight,
                fob_cost,
                amend_fee,
                exw_fee,
                inland_fee,
                total_cost,
                ttl_p_num,
                ttl_p_txt,
                previo,
                insurance,
                str(cell(20) or "").strip(),
            ))

        total_containers = len(master_records)
        total_cost_master = round(sum(x[20] for x in master_records), 2)
        total_previo = round(sum(x[23] for x in master_records), 2)
        total_insurance = round(sum(x[24] for x in master_records), 2)

        acc_balance = 0.0
        acc_pending_receive = 0.0
        acc_pending_note = ""
        acc_total_rcvd = 0.0
        acc_total_cost = 0.0
        acc_paid_mexico = 0.0
        acc_paid_china = 0.0
        payments_records = []
        cust_cost_records = []

        if acc_sheet_name:
            ws_acc = wb[acc_sheet_name]
            acc_rows = list(ws_acc.iter_rows(values_only=True))

            if len(acc_rows) > 0 and len(acc_rows[0]) > 1:
                acc_balance = parse_float(acc_rows[0][1]) or 0.0
            if len(acc_rows) > 1 and len(acc_rows[1]) > 1:
                acc_pending_receive = parse_float(acc_rows[1][1]) or 0.0
                if len(acc_rows[1]) > 2 and acc_rows[1][2]:
                    acc_pending_note = str(acc_rows[1][2]).strip()
            if len(acc_rows) > 2:
                if len(acc_rows[2]) > 1:
                    acc_total_rcvd = parse_float(acc_rows[2][1]) or 0.0
                if len(acc_rows[2]) > 8:
                    acc_total_cost = parse_float(acc_rows[2][8]) or 0.0
            if len(acc_rows) > 3 and len(acc_rows[3]) > 1:
                acc_paid_mexico = parse_float(acc_rows[3][1]) or 0.0

            current_section = "MEXICO"
            sig_counter = {}

            for idx, r in enumerate(acc_rows):
                if not r:
                    continue
                c0 = str(r[0] or "").strip()
                c1 = r[1] if len(r) > 1 else None

                if idx >= 5 and len(r) > 8:
                    cname = str(r[7] or "").strip()
                    ccost = parse_float(r[8])
                    crem = str(r[9] or "").strip() if len(r) > 9 and r[9] is not None else ""
                    if cname and cname.upper() not in ("CUSTOMERNAME", "TOTALCOST") and ccost is not None:
                        cust_cost_records.append((raw_date, idx + 1, cname, ccost, crem))

                if idx < 5:
                    continue

                if "PAID IN CHINA" in c0.upper():
                    acc_paid_china = parse_float(c1) or 0.0
                    current_section = "CHINA"
                    continue

                if "RCVD" in c0.upper() and "DATE" in c0.upper():
                    continue

                rcvd_date = parse_date_str(r[0])
                rcvd_amt = parse_float(c1)
                rcvd_amt_text = str(c1 or "").strip()
                paid_from = str(r[2] or "").strip() if len(r) > 2 and r[2] is not None else ""
                paid_to = str(r[3] or "").strip() if len(r) > 3 and r[3] is not None else ""
                remark = str(r[4] or "").strip() if len(r) > 4 and r[4] is not None else ""

                if not rcvd_date and rcvd_amt is None and not paid_from and not remark:
                    continue

                base_sig = f"{current_section}|{rcvd_date}|{rcvd_amt or 0.0}|{paid_from.upper()}|{remark.upper()}"
                cnt = sig_counter.get(base_sig, 0) + 1
                sig_counter[base_sig] = cnt
                payment_sig = f"{base_sig}#{cnt}"

                payments_records.append((
                    raw_date,
                    current_section,
                    idx + 1,
                    rcvd_date,
                    rcvd_amt or 0.0,
                    rcvd_amt_text,
                    paid_from,
                    paid_to,
                    remark,
                    payment_sig
                ))

        cur.execute("DELETE FROM master_items WHERE raw_date = ?", (raw_date,))
        cur.execute("DELETE FROM accounting_payments WHERE raw_date = ?", (raw_date,))
        cur.execute("DELETE FROM accounting_customer_costs WHERE raw_date = ?", (raw_date,))

        cur.execute("""
            INSERT INTO statements (
                raw_date, stmt_date, filename, file_size, file_mtime, imported_at,
                total_containers, total_cost_master, total_previo, total_insurance,
                acc_balance, acc_pending_receive, acc_pending_note,
                acc_total_rcvd, acc_total_cost, acc_paid_mexico, acc_paid_china, sheets_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(raw_date) DO UPDATE SET
                stmt_date=excluded.stmt_date,
                filename=excluded.filename,
                file_size=excluded.file_size,
                file_mtime=excluded.file_mtime,
                imported_at=excluded.imported_at,
                total_containers=excluded.total_containers,
                total_cost_master=excluded.total_cost_master,
                total_previo=excluded.total_previo,
                total_insurance=excluded.total_insurance,
                acc_balance=excluded.acc_balance,
                acc_pending_receive=excluded.acc_pending_receive,
                acc_pending_note=excluded.acc_pending_note,
                acc_total_rcvd=excluded.acc_total_rcvd,
                acc_total_cost=excluded.acc_total_cost,
                acc_paid_mexico=excluded.acc_paid_mexico,
                acc_paid_china=excluded.acc_paid_china,
                sheets_json=excluded.sheets_json
        """, (
            raw_date,
            format_stmt_date(raw_date),
            filepath.name,
            stat.st_size,
            stat.st_mtime,
            datetime.now().isoformat(timespec="seconds"),
            total_containers,
            total_cost_master,
            total_previo,
            total_insurance,
            acc_balance,
            acc_pending_receive,
            acc_pending_note,
            acc_total_rcvd,
            acc_total_cost,
            acc_paid_mexico,
            acc_paid_china,
            json.dumps(sheetnames, ensure_ascii=False)
        ))

        cur.executemany("""
            INSERT INTO master_items (
                raw_date, row_num, ref_key, ref_no, customer, customer_for, container,
                etd, eta, update_eta, shipping_line, pol, pod,
                bill_amt, bill_amt_text, sea_freight, fob_cost, amend_fee, exw_fee, inland_fee,
                total_cost, ttl_cost_profit, ttl_cost_profit_text, previo, insurance, remark
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, master_records)

        cur.executemany("""
            INSERT INTO accounting_payments (
                raw_date, section, row_num, rcvd_date, rcvd_amt, rcvd_amt_text,
                paid_from, paid_to, remark, payment_sig
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, payments_records)

        cur.executemany("""
            INSERT INTO accounting_customer_costs (
                raw_date, row_num, customer_name, cost, remark
            ) VALUES (?, ?, ?, ?, ?)
        """, cust_cost_records)

        conn.commit()
        return {"raw_date": raw_date, "status": "imported", "containers": total_containers}
    finally:
        wb.close()
        conn.close()
        if tmp_path and tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass


def sync_workspace_statements(force: bool = False) -> List[Dict[str, Any]]:
    results = []
    files = sorted(WORKSPACE_DIR.glob("*STATEMENT*.xlsx"))
    for f in files:
        if f.name.startswith("~$"):
            continue
        try:
            res = parse_and_store_statement(f, force=force)
            results.append(res)
        except Exception as e:
            results.append({"file": f.name, "status": "error", "error": str(e)})
    return results


COST_FIELDS = [
    ("sea_freight", "Sea Freight"),
    ("fob_cost", "FOB Cost"),
    ("amend_fee", "Amend Fee"),
    ("exw_fee", "EXW Fee"),
    ("inland_fee", "Inland Fee"),
    ("previo", "Previo"),
    ("insurance", "Insurance"),
    ("total_cost", "Total Cost"),
    ("bill_amt", "Bill Amt"),
]


def compute_diff(raw_date: str, compare_to: Optional[str] = None) -> Dict[str, Any]:
    conn = get_db()
    cur = conn.cursor()

    all_dates = [r["raw_date"] for r in cur.execute("SELECT raw_date FROM statements ORDER BY raw_date ASC").fetchall()]
    if raw_date not in all_dates:
        conn.close()
        raise HTTPException(status_code=404, detail=f"Statement {raw_date} no encontrado")

    if not compare_to:
        idx = all_dates.index(raw_date)
        compare_to = all_dates[idx - 1] if idx > 0 else None

    curr_stmt = dict(cur.execute("SELECT * FROM statements WHERE raw_date = ?", (raw_date,)).fetchone())
    prev_stmt = dict(cur.execute("SELECT * FROM statements WHERE raw_date = ?", (compare_to,)).fetchone()) if compare_to else None

    curr_items = {r["ref_key"]: dict(r) for r in cur.execute("SELECT * FROM master_items WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()}
    prev_items = {r["ref_key"]: dict(r) for r in cur.execute("SELECT * FROM master_items WHERE raw_date = ? ORDER BY row_num ASC", (compare_to,)).fetchall()} if compare_to else {}

    reviews = {
        r["change_key"]: {"status": r["status"], "note": r["note"], "updated_at": r["updated_at"]}
        for r in cur.execute("SELECT * FROM review_status WHERE raw_date = ?", (raw_date,)).fetchall()
    }

    new_shipments = []
    removed_shipments = []
    cost_changes = []
    eta_changes = []

    if compare_to:
        for ref_key, item in curr_items.items():
            if ref_key not in prev_items:
                ckey = f"{raw_date}:NEW:{item['ref_no']}"
                rev = reviews.get(ckey, {"status": "pending", "note": ""})
                new_shipments.append({
                    **item,
                    "change_key": ckey,
                    "review_status": rev["status"],
                    "review_note": rev["note"],
                })
            else:
                p = prev_items[ref_key]
                field_diffs = []
                extra_charges_added = []

                for f_col, f_label in COST_FIELDS:
                    v_old = p.get(f_col) or 0.0
                    v_new = item.get(f_col) or 0.0
                    if abs(v_new - v_old) >= 0.01:
                        delta = round(v_new - v_old, 2)
                        field_diffs.append({
                            "field": f_col,
                            "label": f_label,
                            "old": round(v_old, 2),
                            "new": round(v_new, 2),
                            "delta": delta,
                        })
                        if f_col in ("amend_fee", "exw_fee", "inland_fee", "previo", "insurance", "fob_cost", "sea_freight"):
                            extra_charges_added.append(f"{f_label}: {'+' if delta > 0 else ''}${delta:,.2f}")

                if field_diffs:
                    ckey = f"{raw_date}:COST:{item['ref_no']}"
                    rev = reviews.get(ckey, {"status": "pending", "note": ""})
                    total_old = (p.get("total_cost") or 0.0) + (p.get("previo") or 0.0) + (p.get("insurance") or 0.0)
                    total_new = (item.get("total_cost") or 0.0) + (item.get("previo") or 0.0) + (item.get("insurance") or 0.0)
                    cost_changes.append({
                        "change_key": ckey,
                        "ref_no": item["ref_no"],
                        "customer": item["customer"],
                        "customer_for": item["customer_for"],
                        "container": item["container"],
                        "shipping_line": item["shipping_line"],
                        "pol": item["pol"],
                        "pod": item["pod"],
                        "eta": item["eta"],
                        "update_eta": item["update_eta"],
                        "old_total_cost": round(p.get("total_cost") or 0.0, 2),
                        "new_total_cost": round(item.get("total_cost") or 0.0, 2),
                        "net_delta": round(total_new - total_old, 2),
                        "diffs": field_diffs,
                        "summary_badges": extra_charges_added,
                        "review_status": rev["status"],
                        "review_note": rev["note"],
                        "item": item,
                    })

                op_diffs = []
                for op_col, op_label in [("eta", "ETA"), ("update_eta", "Update ETA"), ("container", "Container")]:
                    o_val = str(p.get(op_col) or "").strip()
                    n_val = str(item.get(op_col) or "").strip()
                    if o_val != n_val:
                        op_diffs.append({"field": op_col, "label": op_label, "old": o_val, "new": n_val})

                if op_diffs:
                    ckey = f"{raw_date}:ETA:{item['ref_no']}"
                    rev = reviews.get(ckey, {"status": "pending", "note": ""})
                    eta_changes.append({
                        "change_key": ckey,
                        "ref_no": item["ref_no"],
                        "customer": item["customer"],
                        "customer_for": item["customer_for"],
                        "container": item["container"],
                        "diffs": op_diffs,
                        "review_status": rev["status"],
                        "review_note": rev["note"],
                    })

        for ref_key, p_item in prev_items.items():
            if ref_key not in curr_items:
                removed_shipments.append(p_item)

    curr_payments = [dict(r) for r in cur.execute("SELECT * FROM accounting_payments WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()]
    prev_pay_sigs = set()
    if compare_to:
        prev_pay_sigs = {
            r["payment_sig"]
            for r in cur.execute("SELECT payment_sig FROM accounting_payments WHERE raw_date = ?", (compare_to,)).fetchall()
        }

    new_payments = []
    if compare_to:
        for pay in curr_payments:
            if pay["payment_sig"] not in prev_pay_sigs:
                ckey = f"{raw_date}:PAY:{pay['payment_sig']}"
                rev = reviews.get(ckey, {"status": "pending", "note": ""})
                new_payments.append({
                    **pay,
                    "change_key": ckey,
                    "review_status": rev["status"],
                    "review_note": rev["note"],
                })

    curr_cust_costs = [dict(r) for r in cur.execute("SELECT * FROM accounting_customer_costs WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()]
    prev_cust_map = {}
    if compare_to:
        prev_cust_map = {
            r["customer_name"]: r["cost"]
            for r in cur.execute("SELECT customer_name, cost FROM accounting_customer_costs WHERE raw_date = ?", (compare_to,)).fetchall()
        }

    customer_cost_diffs = []
    for cc in curr_cust_costs:
        old_c = prev_cust_map.get(cc["customer_name"], 0.0) if compare_to else cc["cost"]
        delta_c = round(cc["cost"] - old_c, 2)
        customer_cost_diffs.append({
            **cc,
            "old_cost": round(old_c, 2),
            "delta": delta_c,
        })

    conn.close()

    total_alert_items = len(cost_changes) + len(new_shipments) + len(new_payments)
    reviewed_alert_items = (
        sum(1 for x in cost_changes if x["review_status"] == "reviewed")
        + sum(1 for x in new_shipments if x["review_status"] == "reviewed")
        + sum(1 for x in new_payments if x["review_status"] == "reviewed")
    )

    return {
        "raw_date": raw_date,
        "stmt_date": curr_stmt["stmt_date"],
        "compare_to": compare_to,
        "compare_to_date": prev_stmt["stmt_date"] if prev_stmt else None,
        "current_statement": curr_stmt,
        "previous_statement": prev_stmt,
        "summary": {
            "new_shipments_count": len(new_shipments),
            "removed_shipments_count": len(removed_shipments),
            "cost_changes_count": len(cost_changes),
            "eta_changes_count": len(eta_changes),
            "new_payments_count": len(new_payments),
            "new_payments_total": round(sum(p["rcvd_amt"] for p in new_payments), 2),
            "net_extra_charges_delta": round(sum(c["net_delta"] for c in cost_changes), 2),
            "total_alert_items": total_alert_items,
            "reviewed_alert_items": reviewed_alert_items,
            "pending_alert_items": total_alert_items - reviewed_alert_items,
        },
        "cost_changes": cost_changes,
        "new_shipments": new_shipments,
        "removed_shipments": removed_shipments,
        "eta_changes": eta_changes,
        "new_payments": new_payments,
        "customer_cost_diffs": customer_cost_diffs,
    }


def resolve_target_sheet_for_item(item: Dict[str, Any], cust_to_sheet: Dict, cfor_to_sheet: Dict, prefix_to_sheet: Dict) -> str:
    cf = re.sub(r"\s+", " ", (item.get("customer_for") or "").strip().upper())
    cu = re.sub(r"\s+", " ", (item.get("customer") or "").strip().upper())
    ref = (item.get("ref_no") or "").strip().upper()
    pfx = ref.split("-")[0].strip()

    # Specific deterministic rules from CONCENTRADO distribution
    if cf == "KOBI":
        if "ABY-SKM-ANM" in ref:
            return "KOBI"
        if "SKM" in ref:
            return "SKM"
        if "ANM" in ref:
            return "ANM"
        return "KOBI"
    if cf == "ISAMAR":
        if "GLAM" in cu:
            return "GLAM"
        return "AASZ "
    if "ANHERST" in cf:
        if "CATTAN" in cu or pfx in ("GAC", "GABRI"):
            return "CATTAN"
        return "ANHERST"
    if cf in ("YOSEF SENADO", "YOSEF", "HOLLY"):
        if pfx == "YTM":
            return "YTM"
        return "YOSEF"
    if cf in ("SALOMON JASQUI", "SOLOMON JASQUI", "VISIONEVOLUCION SA DE CV"):
        if "LUCK" in cu:
            return "LUCK"
        return "SALOMON JASQUI"
    if cf == "MLO-AB":
        return "AB"
    if cf == "MLO-HHB":
        return "HHB"
    if cf == "DAVID MANN":
        return "MAN"
    if cf == "SHLOMO GALOR":
        return "RED"
    if "DIVENTO" in cf:
        return "ELJA"
    if "MERCURIO" in cf:
        return "MERCURIO"
    if cf == "HUA":
        return "JOSHUA CREDI"
    if "BEPASI" in cf:
        return "DACH"
    if "JERARKYNG" in cf:
        return "JERARKYNG"
    if "GRUPO LUCK" in cf:
        return "LUCK"
    if cf == "JAZD":
        return "JAZ"

    if (cf, cu) in cust_to_sheet and cust_to_sheet[(cf, cu)]:
        return cust_to_sheet[(cf, cu)].most_common(1)[0][0]
    if cf in cfor_to_sheet and cfor_to_sheet[cf]:
        return cfor_to_sheet[cf].most_common(1)[0][0]
    if pfx in prefix_to_sheet and prefix_to_sheet[pfx]:
        return prefix_to_sheet[pfx].most_common(1)[0][0]
    return "SALOMON JASQUI"


def round_up_50(val: float) -> float:
    import math
    if val <= 0:
        return 0.0
    return float(math.ceil(val / 50.0) * 50.0)


def compute_concentrado_target_values_for_sheet(m: Dict[str, Any], sname: str, hdr: List[str]) -> Dict[str, Any]:
    """
    Calcula los valores exactos que corresponden a cada columna de una pestaña específica
    del Concentrado siguiendo la regla de negocio del usuario:
      1) Si el Statement trae BILL AMT (>0), ése se convierte en el TOTAL COST del Concentrado
         (descontando el tren comercial $1,700 si la pestaña tiene columna propia de RAIL y el BILL AMT ya incluía el tren).
      2) Si el Statement tiene vacío el BILL AMT, se calcula como:
         RedondearHaciaArribaAlMultiploDe50( CostoBaseSinColumnasSeparadas + $350 ).
      3) Previo se cobra en $180 ($200 en JAZ) y Seguro en $100; si la hoja tiene una sola columna
         combinada (ej. PREVIO / INSURANCE o INSURANCE en SALOMON JASQUI), van sumados ($180, $100 o $280).
    """
    has_rail_col = any("RAIL" in h for h in hdr)
    has_exw_col = any("EXW" in h or "OTHER COST" in h for h in hdr)
    has_amend_col = any("AMEND" in h or "DETENTION" in h for h in hdr)
    has_sep_prev_ins = (
        any("PREVIO" in h and "INSURANCE" not in h for h in hdr)
        and any("INSURANCE" in h and "PREVIO" not in h for h in hdr)
    )

    stmt_sea = float(m.get("sea_freight") or 0.0)
    stmt_fob = float(m.get("fob_cost") or 0.0)
    stmt_inland = float(m.get("inland_fee") or 0.0)
    stmt_amend = float(m.get("amend_fee") or 0.0)
    stmt_exw = float(m.get("exw_fee") or 0.0)
    stmt_total = float(m.get("total_cost") or 0.0)
    stmt_prev = float(m.get("previo") or 0.0)
    stmt_ins = float(m.get("insurance") or 0.0)
    stmt_bill = parse_float(m.get("bill_amt"))

    # Rail comercial ($1,700 cuando hay Inland Fee en Statement y la hoja tiene columna RAIL)
    sug_rail = 1700.0 if (stmt_inland > 0 and has_rail_col) else 0.0

    # Previo ($180 o $200 en JAZ) y Seguro ($100)
    unit_prev = 200.0 if sname.strip().upper() == "JAZ" else 180.0
    sug_prev = round((stmt_prev / 140.0) * unit_prev, 2) if stmt_prev > 0 else 0.0
    sug_ins = round(stmt_ins, 2) if stmt_ins > 0 else 0.0

    # Cargos extras en columnas dedicadas vs incluidos en el costo base
    sug_exw = round(stmt_exw, 2) if has_exw_col else 0.0
    sug_amend = round(stmt_amend, 2) if has_amend_col else 0.0

    # Costo base de Eliezer que debe quedar cubierto dentro de la columna TOTAL COST de esta hoja
    base_cost_for_tc = stmt_total
    if has_rail_col and stmt_inland > 0:
        base_cost_for_tc -= stmt_inland
    if has_exw_col and stmt_exw > 0:
        base_cost_for_tc -= stmt_exw
    if has_amend_col and stmt_amend > 0:
        base_cost_for_tc -= stmt_amend
    base_cost_for_tc = max(0.0, round(base_cost_for_tc, 2))

    # Regla del usuario para TOTAL COST en Concentrado:
    # - Si Master List trae BILL AMT, se toma ese BILL AMT (restando $1,700 de tren si la hoja separa RAIL y el BILL AMT incluía el tren).
    # - Si Master List trae vacío el BILL AMT, se calcula: ceil((CostoBase + 350) / 50) * 50.
    if stmt_bill is not None and stmt_bill > 0:
        if has_rail_col and stmt_inland > 0 and stmt_bill >= (stmt_sea + stmt_inland):
            sug_tc = round(stmt_bill - sug_rail, 2)
        else:
            sug_tc = round(stmt_bill, 2)
        formula_label = f"Tomado de BILL AMT (${stmt_bill:,.0f})" + (f" − Rail (${sug_rail:,.0f})" if (has_rail_col and stmt_inland > 0 and stmt_bill >= (stmt_sea + stmt_inland)) else "")
    else:
        sug_tc = round_up_50(base_cost_for_tc + 350.0) if base_cost_for_tc > 0 else 0.0
        formula_label = f"Redondeo múltiplo $50 (Costo ${base_cost_for_tc:,.0f} + $350 → ${sug_tc:,.0f})"

    sug_bill_total = round(sug_tc + sug_rail + sug_prev + sug_ins + sug_exw + sug_amend, 2)

    return {
        "suggested_total_cost": sug_tc,
        "suggested_rail": sug_rail,
        "suggested_previo": sug_prev,
        "suggested_insurance": sug_ins,
        "suggested_previo_ins_combined": round(sug_prev + sug_ins, 2),
        "suggested_exw": sug_exw,
        "suggested_amend": sug_amend,
        "suggested_exw_amend": round(stmt_exw + stmt_amend, 2),
        "suggested_bill_amt": sug_bill_total,
        "has_rail_col": has_rail_col,
        "has_sep_prev_ins": has_sep_prev_ins,
        "formula_label": formula_label,
    }


def audit_statement_vs_concentrado(raw_date: str) -> Dict[str, Any]:
    if not CONCENTRADO_PATH.exists():
        raise HTTPException(status_code=404, detail="No se encontró CONCENTRADO ML LOGISTICS ELIEZER 2026.xlsx")

    conn = get_db()
    all_dates = [r["raw_date"] for r in conn.execute("SELECT raw_date FROM statements ORDER BY raw_date ASC").fetchall()]
    if raw_date not in all_dates:
        conn.close()
        raise HTTPException(status_code=404, detail=f"Statement {raw_date} no encontrado")

    idx_date = all_dates.index(raw_date)
    prev_date = all_dates[idx_date - 1] if idx_date > 0 else None

    stmt_row = dict(conn.execute("SELECT * FROM statements WHERE raw_date = ?", (raw_date,)).fetchone())
    ml_items = [dict(r) for r in conn.execute("SELECT * FROM master_items WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()]
    ml_prev_items = (
        {r["ref_no"]: dict(r) for r in conn.execute("SELECT * FROM master_items WHERE raw_date = ?", (prev_date,)).fetchall()}
        if prev_date else {}
    )
    conn.close()

    wb, tmp_path = safe_load_workbook(CONCENTRADO_PATH, data_only=True, read_only=True)
    try:
        sheets = wb.sheetnames
        cuenta_general = []
        cuenta_general_total = 0.0
        cg_map = {}
        if "CUENTA GENERAL" in sheets:
            ws_cg = wb["CUENTA GENERAL"]
            for idx, r in enumerate(ws_cg.iter_rows(values_only=True)):
                if idx == 1 and r and len(r) > 2 and parse_float(r[2]) is not None:
                    cuenta_general_total = parse_float(r[2]) or 0.0
                if idx >= 4 and r and len(r) > 2 and r[1] is not None:
                    cname = str(r[1]).strip()
                    cmonto = parse_float(r[2])
                    if cname and cmonto is not None:
                        cuenta_general.append({"cliente": cname, "monto": cmonto})
                        cg_map[cname.upper()] = cmonto

        conc_by_ref = collections.defaultdict(list)
        sheet_distribution = []
        sheet_headers_map: Dict[str, List[str]] = {}
        cust_to_sheet = collections.defaultdict(collections.Counter)
        cfor_to_sheet = collections.defaultdict(collections.Counter)
        prefix_to_sheet = collections.defaultdict(collections.Counter)

        ml_by_ref = {r["ref_no"]: r for r in ml_items}
        total_conc_billed = 0.0

        for sname in sheets:
            if sname in ("CUENTA GENERAL", "PAGOS", "Hoja1"):
                continue
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            hdr = [str(c).strip().upper() if c is not None else "" for c in rows[1]]
            sheet_headers_map[sname] = hdr
            r1_balance = next((parse_float(c) for c in (rows[0] if rows else []) if parse_float(c) is not None), None)
            if r1_balance is None:
                r1_balance = cg_map.get(sname.strip().upper(), 0.0)
            bill_col_idx = next((i for i, h in enumerate(hdr) if "BILL AMT" in h), -1)
            if bill_col_idx == -1:
                bill_col_idx = next((i for i, h in enumerate(hdr) if "TOTAL COST" in h), -1)

            row_count = 0
            sheet_billed = 0.0
            cfor_set = set()

            for idx in range(2, len(rows)):
                r = rows[idx]
                if not r or r[0] is None or str(r[0]).strip() == "":
                    continue
                ref = str(r[0]).strip()
                if ref.upper().startswith("REF"):
                    continue
                row_count += 1
                row_map = {hdr[i]: parse_float(r[i]) for i in range(min(len(hdr), len(r))) if hdr[i]}
                row_bill = (parse_float(r[bill_col_idx]) or 0.0) if (bill_col_idx != -1 and bill_col_idx < len(r)) else 0.0
                sheet_billed += row_bill
                cont_col = next((i for i, h in enumerate(hdr) if "CONTAINER" in h), None)
                cont_val = str(r[cont_col] or "").strip() if cont_col is not None and cont_col < len(r) else ""
                conc_by_ref[ref].append({
                    "sheet": sname,
                    "row": idx + 1,
                    "hdr": hdr,
                    "map": row_map,
                    "bill_amt": row_bill,
                    "container": cont_val,
                })

                pfx = ref.split("-")[0].strip().upper()
                prefix_to_sheet[pfx][sname] += 1

                if ref in ml_by_ref:
                    m = ml_by_ref[ref]
                    cf = re.sub(r"\s+", " ", m["customer_for"].strip().upper())
                    cu = re.sub(r"\s+", " ", m["customer"].strip().upper())
                    cfor_set.add(cf)
                    cfor_to_sheet[cf][sname] += 1
                    cust_to_sheet[(cf, cu)][sname] += 1

            total_conc_billed += sheet_billed
            sheet_distribution.append({
                "sheet": sname,
                "rows_count": row_count,
                "balance": r1_balance or 0.0,
                "billed_total": round(sheet_billed, 2),
                "headers": [h for h in hdr[:15] if h],
                "mapped_customers_for": sorted(cfor_set),
            })

        missing_in_concentrado = []
        discrepancies = []
        matched_stmt_cost = 0.0
        matched_conc_billed = 0.0
        margin_by_sheet = collections.defaultdict(lambda: {"refs": 0, "stmt_cost": 0.0, "conc_billed": 0.0})

        for m in ml_items:
            ref = m["ref_no"]
            entries = conc_by_ref.get(ref)
            if not entries:
                for cref, clist in conc_by_ref.items():
                    if any(c["container"] == m["container"] and m["container"] != "" for c in clist):
                        entries = clist
                        break

            target_sheet = resolve_target_sheet_for_item(m, cust_to_sheet, cfor_to_sheet, prefix_to_sheet)
            target_hdr = sheet_headers_map.get(target_sheet, ["REF/INV NO", "CUSTOMER", "CUSTOMER FOR", "CONTAINER", "ETD", "ETA", "SHIPPING LINE", "POL", "POD", "TOTAL COST", "RAIL", "PREVIO / INSURANCE", "BILL AMT"])
            sug = compute_concentrado_target_values_for_item_sheet = compute_concentrado_target_values_for_sheet(m, target_sheet, target_hdr)

            if not entries:
                missing_in_concentrado.append({
                    **m,
                    "target_sheet": target_sheet,
                    **sug,
                })
                continue

            primary = entries[0]
            sname = primary["sheet"]
            rnum = primary["row"]
            hdr = primary["hdr"]
            sug_sheet = compute_concentrado_target_values_for_sheet(m, sname, hdr)

            item_stmt_cost = (m["total_cost"] or 0.0) + (m["previo"] or 0.0) + (m["insurance"] or 0.0)
            item_conc_bill = sum((e.get("bill_amt") or 0.0) for e in entries)
            matched_stmt_cost += item_stmt_cost
            matched_conc_billed += item_conc_bill
            margin_by_sheet[sname]["refs"] += 1
            margin_by_sheet[sname]["stmt_cost"] += item_stmt_cost
            margin_by_sheet[sname]["conc_billed"] += item_conc_bill

            act_total = sum((e["map"].get("TOTAL COST") or 0.0) for e in entries)
            if "TOTAL COST" not in hdr:
                act_total = sum((e["map"].get("BILL AMT") or 0.0) for e in entries)

            # 1) Previo & Insurance check (Previo in Stmt 140 -> 180 in Conc, or 200 in JAZ)
            prev_ins_cols = [h for h in hdr if "PREVIO" in h or "INSURANCE" in h]
            if prev_ins_cols:
                exp_prev = sug_sheet["suggested_previo"]
                exp_ins = sug_sheet["suggested_insurance"]
                exp_sum = sug_sheet["suggested_previo_ins_combined"]

                act_sum = sum(
                    (e["map"].get(h) or 0.0)
                    for e in entries
                    for h in e["hdr"]
                    if "PREVIO" in h or "INSURANCE" in h
                )

                if abs(act_sum - exp_sum) >= 1.0:
                    # Check if the $180 Previo was already baked into TOTAL COST (e.g. sea + 400 + 180 = sea + 580)
                    baked_in_total = (act_sum == 0 and exp_prev == 180.0 and abs((act_total - m["sea_freight"]) - 580.0) <= 5.0)
                    target_col = prev_ins_cols[0]
                    if sug_sheet["has_sep_prev_ins"] and exp_prev > 0 and exp_ins == 0:
                        target_col = next((h for h in hdr if "PREVIO" in h), prev_ins_cols[0])
                    elif sug_sheet["has_sep_prev_ins"] and exp_ins > 0 and exp_prev == 0:
                        target_col = next((h for h in hdr if "INSURANCE" in h), prev_ins_cols[0])
                    discrepancies.append({
                        "id": f"PI:{sname}:{rnum}:{ref}",
                        "category": "PREVIO_SEGURO_INCLUIDO_EN_TOTAL" if baked_in_total else "PREVIO_SEGURO",
                        "severity": "info" if baked_in_total else "high",
                        "sheet": sname,
                        "row": rnum,
                        "ref_no": ref,
                        "customer": m["customer"],
                        "customer_for": m["customer_for"],
                        "container": m["container"],
                        "column_name": target_col,
                        "conc_value": round(act_sum, 2),
                        "expected_value": round(exp_sum, 2),
                        "delta": round(exp_sum - act_sum, 2),
                        "description": (
                            f"En Concentrado {target_col}=${act_sum:,.0f} (los $180 de Previo parecen ya sumados dentro de TOTAL COST=${act_total:,.0f})"
                            if baked_in_total else
                            f"En Statement tiene Previo=${m['previo']:,.0f} (=$180 ML) y Seguro=${m['insurance']:,.0f}, pero en Concentrado ({sname}) la columna '{target_col}' tiene ${act_sum:,.0f} en lugar de ${exp_sum:,.0f}"
                        ),
                        "can_autofix": not baked_in_total,
                        "fix_action": {
                            "type": "SET_CELL",
                            "sheet": sname,
                            "row": rnum,
                            "ref_no": ref,
                            "col_header": target_col,
                            "new_value": round(exp_sum, 2),
                        }
                    })

            # 2) Rail / Inland Fee check
            rail_col = next((h for h in hdr if "RAIL" in h), None)
            act_rail = sum((e["map"].get(rail_col) or 0.0) for e in entries) if rail_col else 0.0
            stmt_inland = m["inland_fee"] or 0.0

            if rail_col:
                if stmt_inland == 0 and act_rail > 0:
                    discrepancies.append({
                        "id": f"RAIL_EXTRA:{sname}:{rnum}:{ref}",
                        "category": "TREN_SOBRANTE",
                        "severity": "high",
                        "sheet": sname,
                        "row": rnum,
                        "ref_no": ref,
                        "customer": m["customer"],
                        "customer_for": m["customer_for"],
                        "container": m["container"],
                        "column_name": rail_col,
                        "conc_value": round(act_rail, 2),
                        "expected_value": 0.0,
                        "delta": round(-act_rail, 2),
                        "description": f"En Concentrado ({sname}) tiene cobrado {rail_col}=${act_rail:,.0f}, pero en el Statement INLAND FEE está en $0 (Destino: {m['pod']})",
                        "can_autofix": True,
                        "fix_action": {
                            "type": "SET_CELL",
                            "sheet": sname,
                            "row": rnum,
                            "ref_no": ref,
                            "col_header": rail_col,
                            "new_value": 0.0,
                        }
                    })
                elif stmt_inland > 0 and act_rail == 0 and act_total < (m["sea_freight"] + stmt_inland):
                    discrepancies.append({
                        "id": f"RAIL_MISS:{sname}:{rnum}:{ref}",
                        "category": "TREN_FALTANTE",
                        "severity": "high",
                        "sheet": sname,
                        "row": rnum,
                        "ref_no": ref,
                        "customer": m["customer"],
                        "customer_for": m["customer_for"],
                        "container": m["container"],
                        "column_name": rail_col,
                        "conc_value": 0.0,
                        "expected_value": 1700.0,
                        "delta": 1700.0,
                        "description": f"El Statement tiene INLAND FEE=${stmt_inland:,.0f}, pero en Concentrado ({sname}) la columna {rail_col} está en $0",
                        "can_autofix": True,
                        "fix_action": {
                            "type": "SET_CELL",
                            "sheet": sname,
                            "row": rnum,
                            "ref_no": ref,
                            "col_header": rail_col,
                            "new_value": 1700.0,
                        }
                    })

            # 3) Cargos Extras Nuevos (Amend Fee / EXW Fee) y Costos menores al Statement usando la regla de BILL AMT o Redondeo Múltiplo de $50 (+350)
            act_exw_amend = sum(
                (e["map"].get(h) or 0.0)
                for e in entries
                for h in e["hdr"]
                if "EXW" in h or "AMEND" in h or "OTHER COST" in h or "DETENTION" in h
            )
            extra_col = next((h for h in hdr if "EXW" in h or "AMEND" in h or "OTHER COST" in h or "DETENTION" in h), "TOTAL COST")

            stmt_cost_no_rail = round((m["total_cost"] or 0.0) - (stmt_inland if rail_col else 0.0), 2)
            conc_rev_no_rail = round(act_total + act_exw_amend, 2)
            p_m = ml_prev_items.get(ref)

            if conc_rev_no_rail < stmt_cost_no_rail and sname != "RAMON BETECH":
                exp_target_val = (
                    round(m["amend_fee"] + m["exw_fee"], 2)
                    if (m["amend_fee"] > 0 and extra_col != "TOTAL COST" and act_total >= m["sea_freight"])
                    else sug_sheet["suggested_total_cost"]
                )
                discrepancies.append({
                    "id": f"UNDER:{sname}:{rnum}:{ref}",
                    "category": "COSTO_MENOR_O_CERO",
                    "severity": "high",
                    "sheet": sname,
                    "row": rnum,
                    "ref_no": ref,
                    "customer": m["customer"],
                    "customer_for": m["customer_for"],
                    "container": m["container"],
                    "column_name": extra_col if (m["amend_fee"] > 0 and extra_col != "TOTAL COST" and act_total >= m["sea_freight"]) else "TOTAL COST",
                    "conc_value": conc_rev_no_rail,
                    "expected_value": exp_target_val,
                    "delta": round(exp_target_val - conc_rev_no_rail, 2),
                    "description": f"En Concentrado ({sname}) tiene ${conc_rev_no_rail:,.2f}, menor al costo de Eliezer (${stmt_cost_no_rail:,.2f}). Sugerido según regla de hoja ({sug_sheet['formula_label']}): ${exp_target_val:,.2f}",
                    "can_autofix": True,
                    "fix_action": {
                        "type": "SET_CELL",
                        "sheet": sname,
                        "row": rnum,
                        "ref_no": ref,
                        "col_header": extra_col if (m["amend_fee"] > 0 and extra_col != "TOTAL COST" and act_total >= m["sea_freight"]) else "TOTAL COST",
                        "new_value": exp_target_val,
                    }
                })
            elif p_m and ((m["amend_fee"] - p_m["amend_fee"]) >= 1.0 or (m["exw_fee"] - p_m["exw_fee"]) >= 1.0):
                if act_exw_amend < (m["amend_fee"] + m["exw_fee"] - 5.0):
                    stmt_extra = round(m["amend_fee"] + m["exw_fee"], 2)
                    new_cell_val = stmt_extra if extra_col != "TOTAL COST" else sug_sheet["suggested_total_cost"]
                    discrepancies.append({
                        "id": f"NEW_EXTRA:{sname}:{rnum}:{ref}",
                        "category": "CARGO_EXTRA_RECIENTE",
                        "severity": "high",
                        "sheet": sname,
                        "row": rnum,
                        "ref_no": ref,
                        "customer": m["customer"],
                        "customer_for": m["customer_for"],
                        "container": m["container"],
                        "column_name": extra_col,
                        "conc_value": round(act_exw_amend if extra_col != "TOTAL COST" else act_total, 2),
                        "expected_value": new_cell_val,
                        "delta": round(stmt_extra, 2),
                        "description": f"Cargo extra agregado en este Statement (Amend=${m['amend_fee']:,.2f}, EXW=${m['exw_fee']:,.2f}) no está reflejado en '{extra_col}' de {sname}",
                        "can_autofix": True,
                        "fix_action": {
                            "type": "SET_CELL",
                            "sheet": sname,
                            "row": rnum,
                            "ref_no": ref,
                            "col_header": extra_col,
                            "new_value": new_cell_val,
                        }
                    })

        # Sort discrepancies: high severity first, then by category
        discrepancies.sort(key=lambda x: (0 if x["severity"] == "high" else 1, x["category"], x["sheet"], x["row"]))

        # Enrich sheet_distribution with margin metrics
        for sd in sheet_distribution:
            mdata = margin_by_sheet.get(sd["sheet"], {"refs": 0, "stmt_cost": 0.0, "conc_billed": 0.0})
            sd["stmt_cost_matched"] = round(mdata["stmt_cost"], 2)
            sd["conc_billed_matched"] = round(mdata["conc_billed"], 2)
            sd["gross_margin"] = round(mdata["conc_billed"] - mdata["stmt_cost"], 2)

        debo_eliezer = round(stmt_row.get("acc_balance") or 0.0, 2)
        me_deben_clientes = round(cuenta_general_total or 0.0, 2)
        flujo_neto_cartera = round(me_deben_clientes - debo_eliezer, 2)

        costo_total_eliezer = round(matched_stmt_cost, 2)
        facturado_clientes_2026 = round(matched_conc_billed, 2)
        ganancia_bruta_operativa = round(facturado_clientes_2026 - costo_total_eliezer, 2)
        margen_pct = round((ganancia_bruta_operativa / facturado_clientes_2026) * 100.0, 2) if facturado_clientes_2026 > 0 else 0.0
        ganancia_por_cont = round(ganancia_bruta_operativa / max(len(ml_items), 1), 2)

        return {
            "exists": True,
            "raw_date": raw_date,
            "filename": CONCENTRADO_PATH.name,
            "cuenta_general_total": cuenta_general_total,
            "cuenta_general": cuenta_general,
            "sheet_distribution": sheet_distribution,
            "missing_in_concentrado": missing_in_concentrado,
            "discrepancies": discrepancies,
            "financials": {
                "debo_eliezer": debo_eliezer,
                "eliezer_pending_receive": round(stmt_row.get("acc_pending_receive") or 0.0, 2),
                "eliezer_total_paid": round(stmt_row.get("acc_total_rcvd") or 0.0, 2),
                "eliezer_total_cost_acc": round(stmt_row.get("acc_total_cost") or 0.0, 2),
                "me_deben_clientes": me_deben_clientes,
                "clientes_total_billed_all": round(total_conc_billed, 2),
                "clientes_total_paid_est": round(total_conc_billed - me_deben_clientes, 2),
                "flujo_neto_cartera": flujo_neto_cartera,
                "costo_total_eliezer_2026": costo_total_eliezer,
                "facturado_clientes_2026": facturado_clientes_2026,
                "ganancia_bruta_operativa": ganancia_bruta_operativa,
                "margen_porcentaje": margen_pct,
                "ganancia_promedio_contenedor": ganancia_por_cont,
                "total_containers": len(ml_items),
            },
            "summary": {
                "total_sheets": len(sheet_distribution),
                "missing_refs_count": len(missing_in_concentrado),
                "high_discrepancies_count": sum(1 for d in discrepancies if d["severity"] == "high"),
                "info_discrepancies_count": sum(1 for d in discrepancies if d["severity"] != "high"),
            }
        }
    finally:
        wb.close()
        if tmp_path and tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass


class ConcentradoFixAction(BaseModel):
    type: str  # "SET_CELL" or "INSERT_REF"
    sheet: str
    row: Optional[int] = None
    ref_no: str
    col_header: Optional[str] = None
    new_value: Optional[float] = None
    item: Optional[Dict[str, Any]] = None


class ApplyConcentradoFixesRequest(BaseModel):
    actions: List[ConcentradoFixAction]


@app.on_event("startup")
def on_startup():
    init_db()
    sync_workspace_statements(force=False)


@app.get("/api/statements")
def list_statements():
    conn = get_db()
    cur = conn.cursor()
    rows = [dict(r) for r in cur.execute("SELECT * FROM statements ORDER BY raw_date DESC").fetchall()]
    conn.close()

    asc_rows = list(reversed(rows))
    for i, r in enumerate(asc_rows):
        if i == 0:
            r["delta_containers"] = 0
            r["delta_cost"] = 0.0
            r["delta_rcvd"] = 0.0
            r["prev_raw_date"] = None
        else:
            prev = asc_rows[i - 1]
            r["delta_containers"] = r["total_containers"] - prev["total_containers"]
            r["delta_cost"] = round(r["acc_total_cost"] - prev["acc_total_cost"], 2)
            r["delta_rcvd"] = round(r["acc_total_rcvd"] - prev["acc_total_rcvd"], 2)
            r["prev_raw_date"] = prev["raw_date"]

    return {"statements": rows, "workspace_dir": str(WORKSPACE_DIR)}


@app.post("/api/statements/sync")
def trigger_sync(force: bool = False):
    res = sync_workspace_statements(force=force)
    return {"synced": res, "statements": list_statements()["statements"]}


@app.post("/api/statements/upload")
async def upload_statement(
    file: UploadFile = File(...),
    custom_date: Optional[str] = Form(None)
):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Por favor sube un archivo Excel (.xlsx)")

    raw_date = extract_raw_date_from_filename(file.filename)
    if not raw_date and custom_date:
        cleaned = re.sub(r"[^0-9]", "", custom_date)
        if len(cleaned) == 8:
            raw_date = cleaned

    if not raw_date:
        raise HTTPException(
            status_code=400,
            detail="El nombre del archivo debe incluir la fecha en formato AAAAMMDD (ej. '2026 Eliezer STATEMENT 20261021.xlsx') o indicar la fecha al subirlo."
        )

    target_filename = file.filename
    if raw_date not in target_filename:
        target_filename = f"2026 Eliezer STATEMENT {raw_date}.xlsx"

    target_path = WORKSPACE_DIR / target_filename
    content = await file.read()
    with open(target_path, "wb") as f_out:
        f_out.write(content)

    try:
        info = parse_and_store_statement(target_path, force=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al procesar el Statement: {str(e)}")

    diff_data = compute_diff(raw_date)
    return {
        "uploaded": info,
        "filename": target_filename,
        "raw_date": raw_date,
        "diff": diff_data,
    }


@app.get("/api/statements/{raw_date}")
def get_statement_detail(
    raw_date: str,
    search: Optional[str] = Query(None),
    customer_for: Optional[str] = Query(None),
    only_extra_charges: bool = Query(False),
):
    conn = get_db()
    cur = conn.cursor()
    stmt = cur.execute("SELECT * FROM statements WHERE raw_date = ?", (raw_date,)).fetchone()
    if not stmt:
        conn.close()
        raise HTTPException(status_code=404, detail="Statement no encontrado")

    query = "SELECT * FROM master_items WHERE raw_date = ?"
    params: List[Any] = [raw_date]

    if customer_for:
        query += " AND UPPER(customer_for) = UPPER(?)"
        params.append(customer_for)

    if only_extra_charges:
        query += " AND (amend_fee > 0 OR exw_fee > 0 OR inland_fee > 0 OR previo > 0 OR insurance > 0)"

    if search:
        s_like = f"%{search.strip()}%"
        query += " AND (ref_no LIKE ? OR customer LIKE ? OR customer_for LIKE ? OR container LIKE ? OR pod LIKE ?)"
        params.extend([s_like, s_like, s_like, s_like, s_like])

    query += " ORDER BY row_num ASC"
    master_items = [dict(r) for r in cur.execute(query, params).fetchall()]

    customers_for = [
        r["customer_for"]
        for r in cur.execute(
            "SELECT DISTINCT customer_for FROM master_items WHERE raw_date = ? AND customer_for != '' ORDER BY customer_for ASC",
            (raw_date,)
        ).fetchall()
    ]

    payments = [
        dict(r)
        for r in cur.execute("SELECT * FROM accounting_payments WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()
    ]
    customer_costs = [
        dict(r)
        for r in cur.execute("SELECT * FROM accounting_customer_costs WHERE raw_date = ? ORDER BY row_num ASC", (raw_date,)).fetchall()
    ]
    conn.close()

    return {
        "statement": dict(stmt),
        "master_items": master_items,
        "customers_for": customers_for,
        "payments": payments,
        "customer_costs": customer_costs,
    }


@app.get("/api/statements/{raw_date}/diff")
def get_statement_diff(raw_date: str, compare_to: Optional[str] = Query(None)):
    return compute_diff(raw_date, compare_to=compare_to)


@app.get("/api/references/{ref_no:path}/history")
def get_reference_history(ref_no: str):
    conn = get_db()
    cur = conn.cursor()
    rows = [
        dict(r)
        for r in cur.execute("""
            SELECT m.*, s.stmt_date, s.filename
            FROM master_items m
            JOIN statements s ON s.raw_date = m.raw_date
            WHERE m.ref_no = ?
            ORDER BY m.raw_date ASC
        """, (ref_no,)).fetchall()
    ]
    conn.close()
    return {"ref_no": ref_no, "history": rows}


class ReviewToggleRequest(BaseModel):
    change_key: str
    raw_date: str
    change_type: str
    ref_no: str
    status: str
    note: Optional[str] = ""


class ReviewBulkRequest(BaseModel):
    items: List[ReviewToggleRequest]


@app.post("/api/reviews/toggle")
def toggle_review(req: ReviewToggleRequest):
    conn = get_db()
    cur = conn.cursor()
    now_str = datetime.now().isoformat(timespec="seconds")
    cur.execute("""
        INSERT INTO review_status (change_key, raw_date, change_type, ref_no, status, note, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(change_key) DO UPDATE SET
            status=excluded.status,
            note=excluded.note,
            updated_at=excluded.updated_at
    """, (req.change_key, req.raw_date, req.change_type, req.ref_no, req.status, req.note or "", now_str))
    conn.commit()
    conn.close()
    return {"ok": True, "change_key": req.change_key, "status": req.status, "note": req.note}


@app.post("/api/reviews/bulk")
def bulk_review(req: ReviewBulkRequest):
    conn = get_db()
    cur = conn.cursor()
    now_str = datetime.now().isoformat(timespec="seconds")
    data = [
        (item.change_key, item.raw_date, item.change_type, item.ref_no, item.status, item.note or "", now_str)
        for item in req.items
    ]
    cur.executemany("""
        INSERT INTO review_status (change_key, raw_date, change_type, ref_no, status, note, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(change_key) DO UPDATE SET
            status=excluded.status,
            note=excluded.note,
            updated_at=excluded.updated_at
    """, data)
    conn.commit()
    conn.close()
    return {"ok": True, "updated": len(data)}


@app.get("/api/concentrado/overview")
def get_concentrado_overview(raw_date: Optional[str] = Query(None)):
    conn = get_db()
    if not raw_date:
        r = conn.execute("SELECT raw_date FROM statements ORDER BY raw_date DESC LIMIT 1").fetchone()
        raw_date = r["raw_date"] if r else None
    conn.close()
    if not raw_date:
        return {"exists": False}
    return audit_statement_vs_concentrado(raw_date)


@app.post("/api/concentrado/apply-fixes")
def apply_concentrado_fixes(req: ApplyConcentradoFixesRequest):
    if not CONCENTRADO_PATH.exists():
        raise HTTPException(status_code=404, detail="No se encontró CONCENTRADO ML LOGISTICS ELIEZER 2026.xlsx")

    # Backup before modifying
    backup_dir = BASE_DIR / "backups"
    backup_dir.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(str(CONCENTRADO_PATH), str(backup_dir / f"CONCENTRADO_BACKUP_{ts}.xlsx"))

    try:
        wb = openpyxl.load_workbook(str(CONCENTRADO_PATH))
    except PermissionError:
        raise HTTPException(
            status_code=400,
            detail="El archivo 'CONCENTRADO ML LOGISTICS ELIEZER 2026.xlsx' está abierto en Excel. Ciérralo un momento para aplicar los cambios directamente o descárgalo desde el botón."
        )

    applied_count = 0
    for act in req.actions:
        if act.sheet not in wb.sheetnames:
            continue
        ws = wb[act.sheet]
        hdr_row = [str(ws.cell(row=2, column=c).value or "").strip().upper() for c in range(1, 25)]

        if act.type == "SET_CELL" and act.row and act.col_header:
            target_col_idx = next((i + 1 for i, h in enumerate(hdr_row) if h == act.col_header.strip().upper()), None)
            if target_col_idx:
                ws.cell(row=act.row, column=target_col_idx, value=act.new_value)
                applied_count += 1

        elif act.type == "INSERT_REF" and act.item:
            it = act.item
            # Find first empty row in col A starting from row 3
            next_row = 3
            while ws.cell(row=next_row, column=1).value is not None and str(ws.cell(row=next_row, column=1).value).strip() != "":
                next_row += 1

            has_sep_prev_ins = ("PREVIO" in hdr_row and "INSURANCE" in hdr_row)
            has_sep_exw_amend = (any("EXW" in h for h in hdr_row if h) and any("AMEND" in h for h in hdr_row if h))

            for col_i, h in enumerate(hdr_row, start=1):
                if not h:
                    continue
                if "REF/INV" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("ref_no"))
                elif h == "CUSTOMER":
                    ws.cell(row=next_row, column=col_i, value=it.get("customer"))
                elif h == "CUSTOMER FOR":
                    ws.cell(row=next_row, column=col_i, value=it.get("customer_for"))
                elif h == "CONTAINER":
                    ws.cell(row=next_row, column=col_i, value=it.get("container"))
                elif h == "ETD":
                    ws.cell(row=next_row, column=col_i, value=it.get("etd"))
                elif h == "ETA":
                    ws.cell(row=next_row, column=col_i, value=it.get("update_eta") or it.get("eta"))
                elif h == "SHIPPING LINE":
                    ws.cell(row=next_row, column=col_i, value=it.get("shipping_line"))
                elif h == "POL":
                    ws.cell(row=next_row, column=col_i, value=it.get("pol"))
                elif h == "POD":
                    ws.cell(row=next_row, column=col_i, value=it.get("pod"))
                elif h == "TOTAL COST":
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_total_cost", 0.0))
                elif "RAIL" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_rail", 0.0))
                elif "PREVIO" in h and "INSURANCE" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_previo_ins_comb", (it.get("suggested_previo", 0.0) + it.get("suggested_insurance", 0.0))))
                elif "PREVIO" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_previo", 0.0))
                elif "INSURANCE" in h:
                    if has_sep_prev_ins:
                        ws.cell(row=next_row, column=col_i, value=it.get("suggested_insurance", 0.0))
                    else:
                        # Hoja con columna única llamada INSURANCE (ej. SALOMON JASQUI) donde van sumados Previo + Seguro
                        ws.cell(row=next_row, column=col_i, value=it.get("suggested_previo_ins_comb", (it.get("suggested_previo", 0.0) + it.get("suggested_insurance", 0.0))))
                elif "EXW" in h and "AMEND" not in h and has_sep_exw_amend:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_exw_only", it.get("exw", 0.0)))
                elif "AMEND" in h and "EXW" not in h and has_sep_exw_amend:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_amend_only", it.get("amend", 0.0)))
                elif "EXW" in h or "AMEND" in h or "OTHER COST" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_exw_amend", 0.0))
                elif "DETENTION" in h:
                    ws.cell(row=next_row, column=col_i, value=it.get("suggested_detention", it.get("demurrage", 0.0)))
                elif h == "BILL AMT":
                    # Buscar fórmula en la fila inmediata anterior o en alguna fila superior de la hoja
                    formula_copied = False
                    for search_r in range(next_row - 1, 2, -1):
                        prev_formula = ws.cell(row=search_r, column=col_i).value
                        if isinstance(prev_formula, str) and prev_formula.startswith("="):
                            new_f = re.sub(rf"\b([A-Z]+){search_r}\b", rf"\g<1>{next_row}", prev_formula)
                            ws.cell(row=next_row, column=col_i, value=new_f)
                            formula_copied = True
                            break
                    if not formula_copied:
                        ws.cell(row=next_row, column=col_i, value=it.get("suggested_bill_amt", it.get("suggested_total_cost", 0.0)))
            applied_count += 1

    try:
        wb.save(str(CONCENTRADO_PATH))
    except PermissionError:
        wb.close()
        raise HTTPException(
            status_code=400,
            detail="No se pudo guardar porque 'CONCENTRADO ML LOGISTICS ELIEZER 2026.xlsx' está abierto en Excel. Ciérralo e intenta de nuevo."
        )
    wb.close()
    return {"ok": True, "applied": applied_count}


@app.get("/api/concentrado/download")
def download_concentrado():
    if not CONCENTRADO_PATH.exists():
        raise HTTPException(status_code=404, detail="Archivo Concentrado no encontrado")
    tmp_dl = Path(tempfile.gettempdir()) / f"CONCENTRADO_ML_LOGISTICS_ELIEZER_2026_{datetime.now().strftime('%Y%m%d')}.xlsx"
    shutil.copy2(str(CONCENTRADO_PATH), str(tmp_dl))
    return FileResponse(
        path=str(tmp_dl),
        filename=f"CONCENTRADO ML LOGISTICS ELIEZER 2026 ({datetime.now().strftime('%Y-%m-%d')}).xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


@app.get("/")
def serve_index():
    return FileResponse(str(STATIC_DIR / "index.html"))
